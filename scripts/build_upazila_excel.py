"""
Build an Excel workbook of Bangladesh geography (Division | District | Upazila)
straight from the project's authoritative seed file, database/02_seed_geography.sql.

Parsing strategy (no external SQL engine needed):
  * divisions:  INSERT INTO divisions (name, bn_name, code) VALUES ('Dhaka','..','DHA'),
      -> map code -> english name
  * districts:  INSERT INTO districts (... ) VALUES
                  ((SELECT id FROM divisions WHERE code='DHA'), 'Dhaka','..','DHA-DHA', TRUE),
      -> map district-code -> (division-code, english name)
  * upazilas:   INSERT INTO upazilas (... ) VALUES
                  ((SELECT id FROM districts WHERE code='DHA-DHA'), 'Dhaka Sadar','..','DHA-DHA-01', TRUE),
      -> (district-code, english name)

Rows are emitted in the file's natural order (division -> district -> upazila),
so they read top-to-bottom the way the seed is written.
"""

import re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SQL = ROOT / "database" / "02_seed_geography.sql"
OUT = ROOT / "Bangladesh_Divisions_Districts_Upazilas.xlsx"

text = SQL.read_text(encoding="utf-8")

# --- divisions: code -> english name -------------------------------------
# Rows look like: ('Dhaka', 'ঢাকা', 'DHA'),
div_name_by_code = {}
div_block = re.search(r"INSERT INTO divisions.*?VALUES(.*?);", text, re.S)
if div_block:
    for m in re.finditer(r"\(\s*'([^']*)'\s*,\s*'[^']*'\s*,\s*'([A-Z]{3})'\s*\)", div_block.group(1)):
        eng, code = m.group(1), m.group(2)
        div_name_by_code[code] = eng

# --- districts: district-code -> (division-code, english name) -----------
# Rows: ((SELECT id FROM divisions WHERE code='DHA'), 'Dhaka', 'ঢাকা', 'DHA-DHA', TRUE),
district_info = {}  # code -> (div_code, eng_name)
district_order = []  # keep first-seen order
for m in re.finditer(
    r"\(\s*\(SELECT id FROM divisions WHERE code='([A-Z]{3})'\)\s*,\s*'([^']*)'\s*,\s*'[^']*'\s*,\s*'([A-Z]{3}-[A-Z]{3})'",
    text,
):
    div_code, eng, dcode = m.group(1), m.group(2), m.group(3)
    if dcode not in district_info:
        district_info[dcode] = (div_code, eng)
        district_order.append(dcode)

# --- upazilas: (district-code, english name) in file order ---------------
# Rows: ((SELECT id FROM districts WHERE code='DHA-DHA'), 'Dhaka Sadar', 'ঢাকা সদর', 'DHA-DHA-01', TRUE),
upazilas = []  # list of (div_name, dist_name, upa_name)
for m in re.finditer(
    r"\(\s*\(SELECT id FROM districts WHERE code='([A-Z]{3}-[A-Z]{3})'\)\s*,\s*'([^']*)'\s*,",
    text,
):
    dcode, upa = m.group(1), m.group(2)
    dinfo = district_info.get(dcode)
    if not dinfo:
        continue
    div_code, dist_name = dinfo
    div_name = div_name_by_code.get(div_code, div_code)
    upazilas.append((div_name, dist_name, upa))

# --- sanity summary -------------------------------------------------------
print(f"Divisions: {len(div_name_by_code)}")
print(f"Districts: {len(district_info)}")
print(f"Upazilas : {len(upazilas)}")

# --- build the workbook ---------------------------------------------------
wb = Workbook()
ws = wb.active
ws.title = "Upazilas"

headers = ["#", "Division", "District", "Upazila"]
ws.append(headers)

header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(bold=True, color="FFFFFF", size=11)
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border

# Data rows, zebra striped, with a running number.
stripe = PatternFill("solid", fgColor="F2F6FC")
for i, (div, dist, upa) in enumerate(upazilas, start=1):
    r = i + 1
    ws.cell(row=r, column=1, value=i)
    ws.cell(row=r, column=2, value=div)
    ws.cell(row=r, column=3, value=dist)
    ws.cell(row=r, column=4, value=upa)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=r, column=c)
        cell.border = border
        if i % 2 == 0:
            cell.fill = stripe
    ws.cell(row=r, column=1).alignment = Alignment(horizontal="center")

# Column widths.
widths = [6, 16, 20, 26]
for idx, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(idx)].width = w

# Freeze the header and enable an auto-filter over the data.
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(upazilas) + 1}"

# --- second sheet: a count summary per division/district -----------------
ws2 = wb.create_sheet("Summary")
ws2.append(["Division", "District", "Upazila count"])
for c in range(1, 4):
    cell = ws2.cell(row=1, column=c)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")
    cell.border = border

# counts per district in file order
counts = {}
for div, dist, _ in upazilas:
    counts.setdefault((div, dist), 0)
    counts[(div, dist)] += 1

rr = 2
for dcode in district_order:
    div_code, dist_name = district_info[dcode]
    div_name = div_name_by_code.get(div_code, div_code)
    n = counts.get((div_name, dist_name), 0)
    ws2.cell(row=rr, column=1, value=div_name)
    ws2.cell(row=rr, column=2, value=dist_name)
    ws2.cell(row=rr, column=3, value=n).alignment = Alignment(horizontal="center")
    rr += 1

ws2.cell(row=rr, column=2, value="TOTAL").font = Font(bold=True)
ws2.cell(row=rr, column=3, value=len(upazilas)).font = Font(bold=True)
ws2.cell(row=rr, column=3).alignment = Alignment(horizontal="center")

for idx, w in enumerate([16, 22, 14], start=1):
    ws2.column_dimensions[get_column_letter(idx)].width = w
ws2.freeze_panes = "A2"

wb.save(OUT)
print(f"Saved: {OUT}")
